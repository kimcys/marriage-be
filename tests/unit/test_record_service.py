from __future__ import annotations

from pathlib import Path
from uuid import UUID

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from marriage_ocr_api.db.base import Base
from marriage_ocr_api.db.repositories import create_job
from marriage_ocr_api.jobs.status import JobStatus
from marriage_ocr_api.records.importer import EMPTY_FIELD_PLACEHOLDER, import_records_from_xlsx
from marriage_ocr_api.records.models import RecordRevision
from marriage_ocr_api.records.repositories import RecordNotFoundError, create_record, get_record, list_records
from marriage_ocr_api.records.service import (
    RecordConflictError,
    apply_correction,
    approve_record,
    bulk_approve_records,
    delete_record,
)
from marriage_ocr_api.records.status import RecordStatus


@pytest.fixture
def session() -> Session:
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    db = session_factory()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def _job(session: Session) -> UUID:
    job_id = UUID("123e4567-e89b-12d3-a456-426614175400")
    create_job(
        session,
        id=job_id,
        status=JobStatus.COMPLETED,
        original_filename="register.pdf",
        stored_filename="source.pdf",
        content_type="application/pdf",
        file_size_bytes=123,
        input_relative_path="jobs/123/input/source.pdf",
        output_relative_path="jobs/123/output/result.xlsx",
        debug_relative_path="jobs/123/debug",
        stdout_log_relative_path="jobs/123/logs/stdout.log",
        stderr_log_relative_path="jobs/123/logs/stderr.log",
        ocr_git_ref="abc123",
    )
    session.commit()
    return job_id


def test_import_records_from_xlsx_is_idempotent(session: Session, tmp_path: Path) -> None:
    job_id = _job(session)
    payload = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["full_name", "Confidence", "Review Reason", "Source Page", "Source Record"])
    worksheet.append(["Ada Lovelace", 0.97, "", 1, "record_001"])
    worksheet.append(["Grace Hopper", 0.95, "", 1, "record_002"])
    workbook.save(payload)

    created_first = import_records_from_xlsx(session, job_id, payload)
    created_second = import_records_from_xlsx(session, job_id, payload)

    assert created_first == 2
    assert created_second == 0


def test_import_records_from_xlsx_fills_genuinely_empty_fields_with_placeholder(
    session: Session, tmp_path: Path
) -> None:
    job_id = _job(session)
    payload = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["full_name", "ic_number", "Confidence", "Missing Fields", "Source Record"])
    # ic_number is genuinely blank in the source document -- high confidence,
    # so this isn't an uncertain OCR read, just nothing there to read.
    worksheet.append(["Ada Lovelace", None, 0.95, "ic_number", "record_001"])
    workbook.save(payload)

    import_records_from_xlsx(session, job_id, payload)
    session.commit()

    records = list_records(session, job_id=job_id, batch_id=None, status=None, limit=10, offset=0)
    assert len(records) == 1
    record = records[0]
    # Auto-filled with the placeholder rather than left blank...
    assert record.field_values["ic_number"] == EMPTY_FIELD_PLACEHOLDER
    # ...and specifically because it's no longer "missing" -- confidence is
    # high, so nothing here should force a reviewer to act on it.
    assert record.missing_fields == []
    assert record.status == RecordStatus.APPROVED.value


def test_import_records_from_xlsx_flags_low_confidence_for_review(session: Session, tmp_path: Path) -> None:
    job_id = _job(session)
    payload = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["full_name", "ic_number", "Confidence", "Missing Fields", "Source Record"])
    # Every field has a value, but Gemini itself was unsure about the read.
    worksheet.append(["Ada Lovelace", "S1234567", 0.5, "", "record_001"])
    workbook.save(payload)

    import_records_from_xlsx(session, job_id, payload)
    session.commit()

    records = list_records(session, job_id=job_id, batch_id=None, status=None, limit=10, offset=0)
    assert len(records) == 1
    record = records[0]
    assert record.missing_fields == []
    assert record.status == RecordStatus.PENDING_REVIEW.value


def test_apply_correction_approves_low_confidence_record_once_fields_are_filled(session: Session) -> None:
    """A reviewer editing a record has looked at the scan, which is the review
    low OCR confidence asked for -- so filling in the last missing field
    approves it instead of leaving it stuck in the queue."""
    job_id = _job(session)
    record = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-1",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.5,
        validation_issues=["Gemini uncertain fields: ic_number"],
        missing_fields=["ic_number"],
    )
    session.commit()

    updated = apply_correction(
        session,
        record.id,
        expected_version=1,
        field_values={"ic_number": "S1234567"},
        reviewer="reviewer@example.com",
        note="filled in from the source scan",
    )

    assert updated.missing_fields == []
    assert updated.status == RecordStatus.APPROVED.value


def test_apply_correction_keeps_record_pending_while_a_field_is_still_missing(session: Session) -> None:
    job_id = _job(session)
    record = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-2",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.9,
        validation_issues=[],
        missing_fields=["ic_number", "age"],
    )
    session.commit()

    updated = apply_correction(
        session,
        record.id,
        expected_version=1,
        field_values={"ic_number": "S1234567"},
        reviewer="reviewer@example.com",
        note=None,
    )

    assert updated.missing_fields == ["age"]
    assert updated.status == RecordStatus.PENDING_REVIEW.value


def test_apply_correction_rejects_stale_version(session: Session) -> None:
    job_id = _job(session)
    record = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-1",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.97,
        validation_issues=[],
    )
    session.commit()

    with pytest.raises(RecordConflictError):
        apply_correction(
            session,
            record.id,
            expected_version=2,
            field_values={"full_name": "Ada Byron"},
            reviewer="reviewer@example.com",
            note="corrected surname",
        )


def test_delete_record_removes_it_and_its_revisions(session: Session) -> None:
    job_id = _job(session)
    record = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-1",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.97,
        validation_issues=[],
    )
    session.commit()
    apply_correction(
        session,
        record.id,
        expected_version=1,
        field_values={"full_name": "Ada Byron"},
        reviewer="reviewer@example.com",
        note="corrected surname",
    )
    session.commit()
    assert session.query(RecordRevision).filter_by(record_id=record.id).count() == 1

    delete_record(session, record.id)

    assert get_record(session, record.id) is None
    assert session.query(RecordRevision).filter_by(record_id=record.id).count() == 0


def test_delete_record_raises_for_missing_record(session: Session) -> None:
    with pytest.raises(RecordNotFoundError):
        delete_record(session, UUID("00000000-0000-0000-0000-000000000000"))


def test_approve_record(session: Session) -> None:
    job_id = _job(session)
    first = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-1",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.97,
        validation_issues=[],
    )
    session.commit()

    approved = approve_record(session, first.id, expected_version=1, reviewer="reviewer@example.com")

    assert approved.status == RecordStatus.APPROVED.value


def test_bulk_approve_records(session: Session) -> None:
    job_id = _job(session)
    first = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-1",
        field_values={"full_name": "Ada Lovelace"},
        confidence=0.97,
        validation_issues=[],
    )
    second = create_record(
        session,
        job_id=job_id,
        source_key="page-1-row-2",
        field_values={"full_name": "Grace Hopper"},
        confidence=0.95,
        validation_issues=[],
    )
    session.commit()

    bulk_approved = bulk_approve_records(session, [first.id, second.id], reviewer="reviewer@example.com")

    assert [record.status for record in bulk_approved] == [RecordStatus.APPROVED.value, RecordStatus.APPROVED.value]


def test_import_records_from_csv_keeps_raw_mas_kahwin_out_of_field_values(session: Session, tmp_path: Path) -> None:
    from marriage_ocr_api.records.importer import import_records_from_csv

    job_id = _job(session)
    payload = tmp_path / "typed.csv"
    payload.write_text(
        "Bil,Nama Suami,Mas Kahwin,Source File,Failed Fields,Mas Kahwin Raw\n"
        "35/92,Ahmad,,a.pdf,Mas Kahwin,$.80.00\n"
        "36/92,Ali,RM 80.00,b.pdf,,\n",
        encoding="utf-8-sig",
    )

    import_records_from_csv(session, job_id, payload)
    session.commit()

    by_bil = {r.field_values["Bil"]: r for r in list_records(session, job_id=job_id, batch_id=None, status=None, limit=10, offset=0)}
    assert "Mas Kahwin Raw" not in by_bil["35/92"].field_values
    assert by_bil["35/92"].normalized_data["raw"] == {"Mas Kahwin": "$.80.00"}
    assert by_bil["35/92"].normalized_data["Nama Suami"] == "Ahmad"
    assert "raw" not in by_bil["36/92"].normalized_data
